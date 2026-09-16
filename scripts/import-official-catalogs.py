"""Lossless, reproducible import of the two specified ministerial XLSX releases.
Requires openpyxl. Does not infer clinical rules or repair source values.
Usage: python scripts/import-official-catalogs.py /directory/containing/downloads
"""
from pathlib import Path
import sys,json,hashlib,gzip,base64,re,collections
import openpyxl
root=Path(__file__).resolve().parents[1]
src=Path(sys.argv[1]); sources=[]
for system,name,release,start,cols,codecol,typecol in [
 ('ICD-10-IM','_ElencoSist_ICD10IM_v.Gamma2.2.xlsx','Gamma 2.2',4,20,3,5),
 ('CIPI','_ElencoSist_CIPI_v.Gamma2.1.xlsx','Gamma 2.1',8,13,6,4)]:
 data=(src/name).read_bytes();s=openpyxl.load_workbook(src/name,read_only=True,data_only=True).active
 rows=list(s.iter_rows(min_row=start,max_col=cols,values_only=True))
 assert all(r[codecol] and r[2] and r[typecol] for r in rows), 'Empty identity/type'
 assert all(r[1] in ('Attivo','Storico') for r in rows)
 assert all(re.fullmatch(r'\d{2}/\d{2}/\d{4}\s*-\s*\d{2}/\d{2}/\d{4}',r[-1]) for r in rows)
 ids=[r[2] for r in rows]; assert len(set(ids))==len(ids), 'Repeated source identity'
 rows=[list(r) for r in rows]
 raw=json.dumps(rows,ensure_ascii=False,separators=(',',':')).encode()
 sources.append(dict(system=system,file=name,release=release,url='https://www.salute.gov.it/new/sites/default/files/'+name,sha256=hashlib.sha256(data).hexdigest(),sheet=s.title,firstRow=start,headers=list(next(s.iter_rows(min_row=start-1,max_row=start-1,max_col=cols,values_only=True))),preface=[list(r) for r in s.iter_rows(max_row=start-2,max_col=cols,values_only=True)],rows=rows,rowCount=len(rows),rowsSha256=hashlib.sha256(raw).hexdigest(),statusCounts=dict(collections.Counter(r[1] for r in rows)),terminalCount=sum(r[8]==1 if system=='ICD-10-IM' else r[4]=='Codificante' for r in rows)))
raw=json.dumps(sources,ensure_ascii=False,separators=(',',':')).encode()
compressed=gzip.compress(raw,mtime=0)
out=root/'src/data/official-catalog.payload.json';out.write_text(json.dumps({'encoding':'gzip+base64','payload':base64.b64encode(compressed).decode()},separators=(',',':'))+'\n')
manifest=[{k:v for k,v in s.items() if k not in ('rows','preface','headers')} for s in sources]
(root/'docs/catalog-manifest.json').write_text(json.dumps({'retrieved':'2026-09-16','scope':'Elenco sistematico ICD-10-IM senza capitolo XX; elenco sistematico CIPI. Indice alfabetico non incluso nei due XLSX. Prototipi per sperimentazione.','sources':manifest},ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'sources':manifest,'rawBytes':len(raw),'compressedBytes':len(compressed)},ensure_ascii=False,indent=2))
