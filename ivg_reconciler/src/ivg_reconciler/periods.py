from __future__ import annotations

import csv
import calendar
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class InputPeriod:
    source_file: str
    flow: str
    start: date
    end: date
    label: str


_MONTHS = {
    "GEN": 1, "GENNAIO": 1, "FEB": 2, "FEBBRAIO": 2, "MAR": 3, "MARZO": 3,
    "APR": 4, "APRILE": 4, "MAG": 5, "MAGGIO": 5, "GIU": 6, "GIUGNO": 6,
    "LUG": 7, "LUGLIO": 7, "AGO": 8, "AGOSTO": 8, "SET": 9, "SETTEMBRE": 9,
    "OTT": 10, "OTTOBRE": 10, "NOV": 11, "NOVEMBRE": 11, "DIC": 12, "DICEMBRE": 12,
}


def _quarter_bounds(year: int, q: int) -> tuple[date, date]:
    sm = (q - 1) * 3 + 1
    em = q * 3
    return date(year, sm, 1), date(year, em, calendar.monthrange(year, em)[1])


def _quarter_from_label(label: str) -> int | None:
    s = (label or "").upper()
    if "GENNAIO" in s or "GEN-MAR" in s or "GEN MAR" in s:
        return 1
    if "APRILE" in s or "APR-GIU" in s or "APR GIU" in s:
        return 2
    if "LUGLIO" in s or "LUG-SET" in s or "LUG SET" in s:
        return 3
    if "OTTOBRE" in s or "OTT-DIC" in s or "OTT DIC" in s:
        return 4
    return None


def _detect_csv_delimiter(path: Path) -> str:
    raw = path.read_text(encoding="utf-8-sig", errors="replace")[:4096]
    return ";" if raw.count(";") >= raw.count(",") else ","


def inspect_gino_period(path: str | Path) -> InputPeriod:
    p = Path(path)
    delim = _detect_csv_delimiter(p)
    periods: set[tuple[int, int]] = set()
    flow = ""
    with p.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f, delimiter=delim)
        fields = set(reader.fieldnames or [])
        if "TIPO_IVG" in fields:
            flow = "IVG"
        elif "TIPO_AS" in fields:
            flow = "AS"
        else:
            raise ValueError(f"{p.name}: non riconosciuto come file GINO DOWNLOAD IVG/AS.")
        for row in reader:
            try:
                year = int(float((row.get("ANNO") or "").strip()))
            except Exception:
                try:
                    year = int((row.get("ANNO_ABORTO") or "").strip())
                except Exception:
                    year = 0
            q = _quarter_from_label(row.get("TRIMESTRE") or "")
            if year and q:
                periods.add((year, q))

    if len(periods) == 1:
        year, q = next(iter(periods))
        start, end = _quarter_bounds(year, q)
        return InputPeriod(p.name, flow, start, end, f"Q{q} {year}")

    # Fallback robusto sul nome standard GINO: ..._YYYYMMDD_IVG/AS_DOWNLOAD...
    m = re.search(r"(20\d{6}).*?(IVG|AS).*?DOWNLOAD", p.name.upper())
    if m:
        dt = date(int(m.group(1)[:4]), int(m.group(1)[4:6]), int(m.group(1)[6:8]))
        q = ((dt.month - 1) // 3) + 1
        start, end = _quarter_bounds(dt.year, q)
        return InputPeriod(p.name, flow or m.group(2), start, end, f"Q{q} {dt.year}")

    if not periods:
        raise ValueError(f"{p.name}: periodo GINO non determinabile (TRIMESTRE/ANNO assenti o file vuoto non standard).")
    raise ValueError(f"{p.name}: contiene più periodi dichiarati; impossibile usarlo in modalità fail-proof.")


def _parse_period_text(text: str) -> tuple[date, date] | None:
    s = (text or "").upper().replace("–", "-").replace("—", "-")
    # Intervallo esplicito dd/mm/yyyy - dd/mm/yyyy
    m = re.search(r"(\d{1,2})/(\d{1,2})/(20\d{2})\s*[-A]\s*(\d{1,2})/(\d{1,2})/(20\d{2})", s)
    if m:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1))), date(int(m.group(6)), int(m.group(5)), int(m.group(4)))

    # GEN-GIU 2026, GENNAIO-GIUGNO 2026, ecc.
    month_pattern = "|".join(sorted(_MONTHS, key=len, reverse=True))
    m = re.search(rf"\b({month_pattern})\s*-\s*({month_pattern})\s+(20\d{{2}})\b", s)
    if m:
        y = int(m.group(3)); sm = _MONTHS[m.group(1)]; em = _MONTHS[m.group(2)]
        return date(y, sm, 1), date(y, em, calendar.monthrange(y, em)[1])

    # 1 SEM / 1 SEMESTRE 2026
    m = re.search(r"\b([12])\s*SEM(?:ESTRE)?\s*(20\d{2})\b", s)
    if m:
        sem = int(m.group(1)); y = int(m.group(2))
        return (date(y,1,1), date(y,6,30)) if sem == 1 else (date(y,7,1), date(y,12,31))
    return None


def inspect_sdo_period(path: str | Path) -> InputPeriod:
    p = Path(path)
    wb = load_workbook(p, read_only=True, data_only=True)
    texts: list[str] = []
    flow = ""
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10), values_only=True):
            for v in row[:12]:
                if v not in (None, ""):
                    texts.append(str(v))
    joined = " | ".join(texts)
    u = joined.upper() + " " + p.name.upper()
    if "ABORTI SPONTANEI" in u or re.search(r"\bAS\b", p.name.upper()):
        flow = "AS"
    elif "IVG" in u:
        flow = "IVG"
    else:
        raise ValueError(f"{p.name}: flusso SDO IVG/AS non determinabile.")

    parsed = _parse_period_text(joined) or _parse_period_text(p.name)
    if not parsed:
        raise ValueError(
            f"{p.name}: periodo SDO non determinabile. È richiesto nel report XLSX un testo come "
            "'Accettazioni/Dimissioni nel periodo: GEN-GIU 2026' oppure un nome file con '1 sem 2026'."
        )
    start, end = parsed
    return InputPeriod(p.name, flow, start, end, f"{start.strftime('%d/%m/%Y')}–{end.strftime('%d/%m/%Y')}")


def _merge_contiguous(periods: list[InputPeriod]) -> tuple[date, date]:
    if not periods:
        raise ValueError("Nessun periodo disponibile.")
    ordered = sorted(periods, key=lambda x: (x.start, x.end))
    start, end = ordered[0].start, ordered[0].end
    for p in ordered[1:]:
        if p.start > end.fromordinal(end.toordinal() + 1):
            raise ValueError(
                f"periodi non contigui: {start.strftime('%d/%m/%Y')}–{end.strftime('%d/%m/%Y')} "
                f"e {p.start.strftime('%d/%m/%Y')}–{p.end.strftime('%d/%m/%Y')}"
            )
        if p.end > end:
            end = p.end
    return start, end


def validate_period_files(gino_files: list[str | Path], sdo_files: list[str | Path]) -> tuple[date, date, str]:
    gmeta = [inspect_gino_period(p) for p in gino_files]
    smeta = [inspect_sdo_period(p) for p in sdo_files]
    flows = sorted(set(x.flow for x in gmeta) | set(x.flow for x in smeta))
    errors: list[str] = []
    resolved: dict[str, tuple[date, date]] = {}

    for flow in flows:
        gp = [x for x in gmeta if x.flow == flow]
        sp = [x for x in smeta if x.flow == flow]
        if not gp:
            errors.append(f"{flow}: manca il file GINO per il periodo caricato.")
            continue
        if not sp:
            errors.append(f"{flow}: manca il file SDO XLSX per il periodo caricato.")
            continue
        try:
            gs, ge = _merge_contiguous(gp)
        except ValueError as e:
            errors.append(f"{flow} GINO: {e}")
            continue
        try:
            ss, se = _merge_contiguous(sp)
        except ValueError as e:
            errors.append(f"{flow} SDO: {e}")
            continue
        if (gs, ge) != (ss, se):
            errors.append(
                f"{flow}: periodo GINO {gs.strftime('%d/%m/%Y')}–{ge.strftime('%d/%m/%Y')} "
                f"diverso dal periodo SDO {ss.strftime('%d/%m/%Y')}–{se.strftime('%d/%m/%Y')}."
            )
        else:
            resolved[flow] = (gs, ge)

    if not errors and len(set(resolved.values())) > 1:
        detail = "; ".join(f"{f} {a.strftime('%d/%m/%Y')}–{b.strftime('%d/%m/%Y')}" for f,(a,b) in resolved.items())
        errors.append(f"I flussi caricati non coprono lo stesso periodo complessivo: {detail}.")

    if errors:
        raise ValueError("CONTROLLO PRELIMINARE PERIODI NON SUPERATO:\n\n- " + "\n- ".join(errors) + "\n\nIl matching non è stato eseguito.")

    start, end = next(iter(resolved.values()))
    detail = " | ".join(
        f"{flow}: GINO={start.strftime('%d/%m/%Y')}–{end.strftime('%d/%m/%Y')}; "
        f"SDO={start.strftime('%d/%m/%Y')}–{end.strftime('%d/%m/%Y')}"
        for flow in sorted(resolved)
    )
    return start, end, f"OK – {detail}"
