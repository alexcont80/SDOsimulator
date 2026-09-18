from __future__ import annotations

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .models import ReconciliationResult


def export_report(result: ReconciliationResult, path: str | Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "SCHEDE_MANCANTI"
    missing_headers = [
        "FLUSSO", "NOSOLOGICO_SDO", "DATA_NASCITA", "RESIDENZA", "DATA_AMMISSIONE", "DATA_DIMISSIONE",
        "DRG", "DIAGNOSI", "PROCEDURE", "EPISODIO_DA_AGGIUNGERE", "INDICAZIONE_PUNTUALE"
    ]
    ws.append(missing_headers)
    for m in result.missing:
        ws.append([m.flow, m.sdo_nosologico, m.dob, m.residence, m.admission, m.discharge, m.drg, m.diagnoses, m.procedures, m.episode, m.instruction])

    ws2 = wb.create_sheet("SCHEDE_DA_CORREGGERE")
    corr_headers = [
        "FLUSSO", "ID_GINO", "NOSOLOGICO_SDO", "ITEM_DA_VERIFICARE", "VALORE_GINO", "VALORE_SDO",
        "EPISODIO", "INDICAZIONE_PUNTUALE", "PRIORITA"
    ]
    ws2.append(corr_headers)
    for c in result.corrections:
        ws2.append([c.flow, c.gino_id, c.sdo_nosologico, c.item, c.gino_value, c.sdo_value, c.episode, c.instruction, c.severity])

    ws3 = wb.create_sheet("AUDIT_TECNICO")
    ws3.append(["CATEGORIA", "FLUSSO", "IDENTIFICATIVO", "DETTAGLIO"])
    for a in result.audit:
        ws3.append([a.category, a.flow, a.identifier, a.detail])

    ws4 = wb.create_sheet("SINTESI")
    ws4.append(["INDICATORE", "VALORE"])
    ws4.append(["Record GINO letti", result.gino_input_count])
    ws4.append(["Record GINO validi/deduplicati", result.gino_count])
    ws4.append(["Record SDO nel perimetro", result.sdo_count])
    ws4.append(["Match primari effettuati", result.matched_count])
    ws4.append(["Schede mancanti", len(result.missing)])
    ws4.append(["Schede/item da correggere", len(result.corrections)])

    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in (ws, ws2, ws3, ws4):
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            letter = col[0].column_letter
            max_len = max(len(str(c.value or "")) for c in col[: min(sheet.max_row, 100)])
            sheet.column_dimensions[letter].width = min(max(max_len + 2, 12), 55)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(Path(path))
