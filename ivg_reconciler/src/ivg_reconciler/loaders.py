from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Iterable, Optional

from openpyxl import load_workbook

from .models import GinoRecord, SdoRecord, AuditItem, normalize_text, parse_intish, parse_excel_date, parse_gino_date, parse_gino_dob, norm_code


def load_municipality_map(path: Optional[str | Path]) -> dict[str, str]:
    if not path:
        return {}
    p = Path(path)
    if not p.exists():
        return {}
    with p.open("r", encoding="utf-8-sig", newline="") as f:
        sample = f.read(2048)
        f.seek(0)
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        reader = csv.DictReader(f, delimiter=delimiter)
        result: dict[str, str] = {}
        for row in reader:
            code = (row.get("code") or row.get("CODICE") or row.get("istat_code") or "").strip()
            name = row.get("name") or row.get("COMUNE") or row.get("municipality") or ""
            if code and name:
                result[code.zfill(6)] = normalize_text(name)
        return result


def _detect_delimiter(path: Path) -> str:
    raw = path.read_text(encoding="utf-8-sig", errors="replace")[:4096]
    return ";" if raw.count(";") >= raw.count(",") else ","


def load_gino_csv(path: str | Path, municipality_map: dict[str, str]) -> tuple[list[GinoRecord], list[AuditItem]]:
    p = Path(path)
    delim = _detect_delimiter(p)
    audits: list[AuditItem] = []
    out: list[GinoRecord] = []
    with p.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f, delimiter=delim)
        fields = set(reader.fieldnames or [])
        if "TIPO_IVG" in fields:
            flow = "IVG"
        elif "TIPO_AS" in fields:
            flow = "AS"
        else:
            raise ValueError(f"{p.name}: impossibile riconoscere il flusso GINO (manca TIPO_IVG/TIPO_AS).")
        required = {
            "PROGRESSIVO_QUEST", "DATA_NASCITA_GG", "DATA_NASCITA_MM", "DATA_NASCITA_AAAA",
            "GIORNO_ABORTO", "MESE_ABORTO", "ANNO_ABORTO", "TRIMESTRE"
        }
        missing_fields = sorted(required - fields)
        if missing_fields:
            raise ValueError(f"{p.name}: colonne GINO mancanti: {', '.join(missing_fields)}")

        for line_no, row in enumerate(reader, start=2):
            dob = parse_gino_dob(row)
            event = parse_gino_date(row)
            if not dob or not event:
                audits.append(AuditItem("RECORD_GINO_SCARTATO", flow, f"{p.name}:{line_no}", "DOB o data evento non valida"))
                continue
            prov = norm_code(row.get("PROV_RESIDENZA_GESTANTE"))
            com = norm_code(row.get("COM_RESIDENZA_GESTANTE"))
            foreign = norm_code(row.get("STATO_ESTERO_RESIDENZA"))
            if foreign and foreign != "000":
                residence_code = f"ESTERO:{foreign}"
                residence_name = None
                residence_comparable = False
            else:
                residence_code = f"{prov}{com}" if prov and com else ""
                residence_name = municipality_map.get(residence_code) if residence_code else None
                residence_comparable = residence_name is not None
            institute = (row.get("COD_ISTITUTO") or row.get("CODICE_ISTITUTO_NEW") or row.get("CODICE_ISTITUTO") or "").strip()
            year = parse_intish(row.get("ANNO")) or event.year
            type_col = "TIPO_IVG" if flow == "IVG" else "TIPO_AS"
            out.append(GinoRecord(
                flow=flow,
                source_file=p.name,
                progressivo=(row.get("PROGRESSIVO_QUEST") or "").strip(),
                institute=institute,
                year=year,
                quarter_label=(row.get("TRIMESTRE") or "").strip(),
                dob=dob,
                event_date=event,
                residence_code=residence_code,
                residence_name=residence_name,
                residence_comparable=residence_comparable,
                type_code=(row.get(type_col) or "").strip(),
                regime=(row.get("REGIME_RICOVERO") or "").strip(),
                access_days=(row.get("GIORNATE_ACCESSI") or "").strip(),
                complication_5=(row.get("COMPLICAZIONE_IVG_5") or "").strip() if flow == "IVG" else "",
                raw=row,
            ))
    return out, audits


def _find_header_row(ws) -> int:
    for r in range(1, min(ws.max_row, 30) + 1):
        vals = [str(ws.cell(r, c).value or "").strip() for c in range(1, min(ws.max_column, 60) + 1)]
        if any(v == "m0 Anno acc." for v in vals):
            return r
    raise ValueError(f"Foglio {ws.title}: intestazione 'm0 Anno acc.' non trovata")


def _flow_from_diagnoses(codes: Iterable[str]) -> Optional[str]:
    codes = [str(x).replace(".", "").strip() for x in codes if str(x).strip()]
    if any(c.startswith("635") for c in codes):
        return "IVG"
    if any(c == "632" or c.startswith("634") for c in codes):
        return "AS"
    return None


def load_sdo_xlsx(path: str | Path) -> tuple[list[SdoRecord], list[AuditItem]]:
    p = Path(path)
    wb = load_workbook(p, read_only=True, data_only=True)
    audits: list[AuditItem] = []
    out: list[SdoRecord] = []
    for ws in wb.worksheets:
        try:
            hr = _find_header_row(ws)
        except ValueError:
            continue
        headers = [ws.cell(hr, c).value for c in range(1, ws.max_column + 1)]
        hmap = {str(h).strip(): idx for idx, h in enumerate(headers) if h not in (None, "")}
        req = ["m0 Progressivo num.nosologico", "m0 Data Nascita", "m0 Data di accettazione", "m0 Data di dimissione", "m0 DRG cod."]
        absent = [h for h in req if h not in hmap]
        if absent:
            raise ValueError(f"{p.name}: colonne SDO mancanti: {', '.join(absent)}")

        for row in ws.iter_rows(min_row=hr + 1, values_only=True):
            if not any(v not in (None, "") for v in row):
                continue
            d = {str(headers[i]).strip(): row[i] for i in range(min(len(headers), len(row))) if headers[i] not in (None, "")}
            diagnoses: list[str] = []
            diagnosis_desc: list[str] = []
            procedures: list[str] = []
            procedure_dates: list[date] = []
            procedure_desc: list[str] = []
            for k, v in d.items():
                if re.fullmatch(r"Diagnosi \d+", k) and v not in (None, ""):
                    diagnoses.append(str(v).replace(".", "").strip())
                elif re.fullmatch(r"Diagnosi \d+ desc\.", k) and v not in (None, ""):
                    diagnosis_desc.append(str(v).strip())
                elif re.fullmatch(r"Intervento \d+", k) and v not in (None, ""):
                    procedures.append(str(v).replace(".", "").strip())
                elif re.fullmatch(r"Intervento \d+ data", k) and v not in (None, ""):
                    pd = parse_excel_date(v)
                    if pd:
                        procedure_dates.append(pd)
                elif re.fullmatch(r"Intervento \d+ desc\.", k) and v not in (None, ""):
                    procedure_desc.append(str(v).strip())
            flow = _flow_from_diagnoses(diagnoses)
            if not flow:
                continue
            dob = parse_excel_date(d.get("m0 Data Nascita"))
            adm = parse_excel_date(d.get("m0 Data di accettazione"))
            dis = parse_excel_date(d.get("m0 Data di dimissione"))
            if not (dob and adm and dis):
                audits.append(AuditItem("RECORD_SDO_SCARTATO", flow, str(d.get("m0 Progressivo num.nosologico")), "Data nascita/ammissione/dimissione non valida"))
                continue
            res_name = None
            res_state = None
            for key in d:
                nk = normalize_text(key)
                if "COMUNE DI RESIDENZA" in nk:
                    res_name = normalize_text(d.get(key)) if d.get(key) else None
                if "STATO DI RESIDENZA" in nk:
                    res_state = normalize_text(d.get(key)) if d.get(key) else None
            out.append(SdoRecord(
                flow=flow,
                source_file=p.name,
                nosologico=str(d.get("m0 Progressivo num.nosologico") or "").strip(),
                state=str(d.get("m0 Stato SDO desc.") or "").strip(),
                facility=str(d.get("m0 Presidio acc. sede desc.") or "").strip(),
                ward=str(d.get("m0 Reparto acc.") or "").strip(),
                dob=dob,
                admission=adm,
                discharge=dis,
                drg=str(d.get("m0 DRG cod.") or "").strip(),
                diagnoses=diagnoses,
                diagnosis_desc=diagnosis_desc,
                procedures=procedures,
                procedure_dates=procedure_dates,
                procedure_desc=procedure_desc,
                residence_name=res_name,
                residence_state=res_state,
                raw=d,
            ))
    if not out:
        raise ValueError(f"{p.name}: nessun record SDO IVG/AS riconosciuto")
    return out, audits
