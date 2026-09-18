from __future__ import annotations

import csv
import re
import sys
from pathlib import Path
from openpyxl import load_workbook


def norm(s):
    return re.sub(r"\s+", " ", str(s or "").strip())


def find_columns(ws):
    for r in range(1, min(ws.max_row, 30) + 1):
        vals=[norm(ws.cell(r,c).value) for c in range(1, ws.max_column+1)]
        low=[v.lower() for v in vals]
        code_idx=None; name_idx=None
        for i,v in enumerate(low):
            if "codice comune" in v and ("alfanumerico" in v or "numerico" in v or "istat" in v):
                code_idx=i
                break
        if code_idx is None:
            for i,v in enumerate(low):
                if "codice comune" in v:
                    code_idx=i; break
        for i,v in enumerate(low):
            if "denominazione" in v and ("italiano" in v or "comune" in v or "italiana" in v):
                name_idx=i; break
        if name_idx is None:
            for i,v in enumerate(low):
                if "denominazione" in v:
                    name_idx=i; break
        if code_idx is not None and name_idx is not None:
            return r, code_idx+1, name_idx+1
    return None


def main(src: str, dest: str):
    wb=load_workbook(src, read_only=True, data_only=True)
    mapping={}
    for ws in wb.worksheets:
        found=find_columns(ws)
        if not found:
            continue
        hr, cc, nc=found
        for row in ws.iter_rows(min_row=hr+1, values_only=True):
            code=row[cc-1] if cc-1 < len(row) else None
            name=row[nc-1] if nc-1 < len(row) else None
            if code in (None,"") or name in (None,""):
                continue
            c=str(code).strip()
            if c.endswith(".0"):
                c=c[:-2]
            digits=re.sub(r"\D", "", c)
            if not digits:
                continue
            c=digits.zfill(6)
            if len(c)!=6:
                continue
            mapping[c]=str(name).strip()
    if len(mapping) < 7000:
        raise SystemExit(f"Mapping comuni insufficiente: {len(mapping)} righe (attese >7000)")
    Path(dest).parent.mkdir(parents=True, exist_ok=True)
    with open(dest,"w",encoding="utf-8-sig",newline="") as f:
        w=csv.writer(f,delimiter=";")
        w.writerow(["code","name"])
        for code,name in sorted(mapping.items()):
            w.writerow([code,name])
    print(f"Creato {dest}: {len(mapping)} comuni")


if __name__=="__main__":
    if len(sys.argv)!=3:
        raise SystemExit("Uso: build_comuni.py <Elenco-comuni-italiani.xlsx> <comuni_istat.csv>")
    main(sys.argv[1],sys.argv[2])
